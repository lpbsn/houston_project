from __future__ import annotations

import io
import math
from datetime import date, datetime
from decimal import Decimal

from houston.brain.bounds import ExtractionBudget, zip_entry_names
from houston.brain.errors import ExtractionRejected
from houston.brain.grid import extract_grid
from houston.brain.numbers import decimal_text
from houston.brain.records import (
    SOURCE_ABSENT,
    SOURCE_EMPTY,
    SOURCE_ERROR,
    SOURCE_EXCEL_NUMBER,
    SOURCE_TEXT,
    ExtractionResult,
    GridCell,
)

_INVALID = "Le fichier XLSX n'est pas exploitable."
_ENCRYPTED = "Ce classeur est chiffré. Exportez un XLSX sans mot de passe."
_CHARTS_ONLY = (
    "Ce classeur ne contient pas de valeurs de cellules exploitables. "
    "Les graphiques ne sont pas lus. Exportez les valeurs en XLSX ou en CSV."
)


def extract_xlsx(raw: bytes, budget: ExtractionBudget) -> ExtractionResult:
    budget.check_time()
    names = zip_entry_names(
        raw,
        budget,
        invalid_code="invalid_xlsx",
        invalid_message=_INVALID,
    )
    warnings = _package_warnings(names)
    formulas, external = _formula_index(raw, budget)
    if external:
        warnings.append(
            {
                "code": "external_link_ignored",
                "message": (
                    "Les liens externes ne sont pas ouverts. "
                    "Une formule sans valeur en cache reste absente."
                ),
            }
        )
    sheet_rows = _sheet_rows(raw, budget, formulas)
    measures = []
    blocks = []
    for title, rows in sheet_rows:
        found_measures, found_blocks = extract_grid(rows, locate=_locate(title, rows))
        measures.extend(found_measures)
        blocks.extend(found_blocks)
    result = ExtractionResult(
        source_format="xlsx",
        measures=measures,
        qualitative=blocks,
        warnings=warnings,
    )
    if not result.analyzable() and any(name.startswith("xl/charts/") for name in names):
        raise ExtractionRejected(code="xlsx_values_unavailable", message=_CHARTS_ONLY)
    return result


def _package_warnings(names: list[str]) -> list[dict]:
    warnings = []
    if any(name.endswith("vbaProject.bin") for name in names):
        warnings.append(
            {
                "code": "macros_not_executed",
                "message": "Les macros du classeur ne sont pas exécutées.",
            }
        )
    if any(name.startswith("xl/charts/") for name in names):
        warnings.append(
            {
                "code": "charts_not_interpreted",
                "message": (
                    "Les graphiques ne sont pas interprétés. "
                    "Seules les cellules sont extraites."
                ),
            }
        )
    if any(name.startswith("xl/comments") for name in names):
        warnings.append(
            {
                "code": "comments_not_extracted",
                "message": "Les commentaires du classeur ne sont pas extraits.",
            }
        )
    return warnings


def _formula_index(raw: bytes, budget: ExtractionBudget) -> tuple[set[tuple[str, int, int]], bool]:
    book = _open(raw, data_only=False)
    formulas: set[tuple[str, int, int]] = set()
    external = False
    try:
        budget.check_count(
            len(book.sheetnames),
            limit=budget.max_sheets,
            code="extraction_sheet_limit",
            label="feuilles",
        )
        for sheet in book.worksheets:
            budget.check_time()
            for cells in sheet.iter_rows():
                budget.check_time()
                budget.check_columns(len(cells))
                for cell in cells:
                    if getattr(cell, "data_type", None) != "f":
                        continue
                    formulas.add((sheet.title, cell.row, cell.column))
                    formula = cell.value
                    formula_text = formula.text if hasattr(formula, "text") else str(formula or "")
                    if "[" in formula_text:
                        external = True
    finally:
        book.close()
    return formulas, external


def _sheet_rows(
    raw: bytes,
    budget: ExtractionBudget,
    formulas: set[tuple[str, int, int]],
) -> list[tuple[str, list[list[GridCell]]]]:
    book = _open(raw, data_only=True)
    sheets: list[tuple[str, list[list[GridCell]]]] = []
    try:
        for sheet in book.worksheets:
            budget.check_time()
            rows: list[list[GridCell]] = []
            for cells in sheet.iter_rows():
                budget.add_row()
                budget.check_columns(len(cells))
                grid_row: list[GridCell] = []
                for cell in cells:
                    budget.add_cell()
                    grid_row.append(_grid_cell(sheet.title, cell, formulas))
                rows.append(grid_row)
            sheets.append((sheet.title, rows))
    finally:
        book.close()
    return sheets


def _grid_cell(sheet_title: str, cell, formulas: set[tuple[str, int, int]]) -> GridCell:
    coordinate = getattr(cell, "coordinate", None)
    base = {"row": cell.row, "column": cell.column, "coordinate": coordinate}
    if (sheet_title, cell.row, cell.column) in formulas:
        if getattr(cell, "data_type", None) == "e" or _is_error_text(cell.value):
            return GridCell(source=SOURCE_ERROR, raw=str(cell.value), **base)
        number = _excel_decimal(cell.value)
        if number is not None:
            return GridCell(
                source=SOURCE_EXCEL_NUMBER,
                number=number,
                raw=decimal_text(number),
                **base,
            )
        return GridCell(source=SOURCE_ABSENT, **base)
    if cell.value is None:
        return GridCell(source=SOURCE_EMPTY, **base)
    if getattr(cell, "data_type", None) == "e" or _is_error_text(cell.value):
        return GridCell(source=SOURCE_ERROR, raw=str(cell.value), **base)
    if isinstance(cell.value, datetime):
        return GridCell(source=SOURCE_TEXT, raw=cell.value.isoformat(sep=" "), **base)
    if isinstance(cell.value, date):
        return GridCell(source=SOURCE_TEXT, raw=cell.value.isoformat(), **base)
    if isinstance(cell.value, bool):
        return GridCell(source=SOURCE_TEXT, raw="TRUE" if cell.value else "FALSE", **base)
    number = _excel_decimal(cell.value)
    if number is not None:
        return GridCell(
            source=SOURCE_EXCEL_NUMBER,
            number=number,
            raw=decimal_text(number),
            **base,
        )
    return GridCell(source=SOURCE_TEXT, raw=str(cell.value), **base)


def _excel_decimal(value) -> Decimal | None:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        return None
    if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
        return None
    if isinstance(value, int):
        return Decimal(value)
    return Decimal(str(value))


def _is_error_text(value) -> bool:
    return isinstance(value, str) and value.upper() in {
        "#DIV/0!",
        "#N/A",
        "#VALUE!",
        "#REF!",
        "#NAME?",
        "#NUM!",
        "#NULL!",
    }


def _locate(title: str, rows: list[list[GridCell]]):
    def locate(row_index: int, column_index: int, header: str | None) -> dict:
        cell = rows[row_index][column_index]
        locator = {
            "format": "xlsx",
            "sheet": title,
            "row": cell.row,
            "column": cell.column,
            "cell": cell.coordinate,
        }
        if header:
            locator["header"] = header
        return locator

    return locate


def _open(raw: bytes, *, data_only: bool):
    from openpyxl import load_workbook

    try:
        return load_workbook(
            io.BytesIO(raw),
            read_only=True,
            data_only=data_only,
            keep_links=False,
        )
    except Exception as exc:
        message = _ENCRYPTED if _looks_encrypted(exc) else _INVALID
        code = "xlsx_encrypted" if _looks_encrypted(exc) else "invalid_xlsx"
        raise ExtractionRejected(code=code, message=message) from exc


def _looks_encrypted(exc: BaseException) -> bool:
    text = str(exc).lower()
    return "password" in text or "encrypt" in text
