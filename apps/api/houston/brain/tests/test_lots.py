"""Section grouping and subdivision. No provider call."""

from __future__ import annotations

import io

import pytest

from houston.brain.configuration import (
    canonical_json,
    digest,
    estimate_tokens,
    provider_fixed_input_tokens,
    useful_input_tokens,
)
from houston.brain.errors import BrainWorkError
from houston.brain.extract import extract_document
from houston.brain.lots import (
    _merge,
    _ReadingMeter,
    plan_lots,
    reading_envelope,
    sections_from_payload,
)


def _paragraph(text: str) -> dict:
    return {"kind": "paragraph", "text": text, "locator": {"block": text[:12]}}


def _table(rows: list[str]) -> dict:
    return {
        "kind": "table",
        "header": "Mois",
        "unit": "EUR",
        "locator": {"sheet": "Suivi"},
        "rows": [{"text": row, "locator": {"row": index + 1}} for index, row in enumerate(rows)],
    }


def test_three_short_sections_form_one_lot():
    sections = [_paragraph("un"), _paragraph("deux"), _paragraph("trois")]
    planned = plan_lots(
        sections,
        per_call_input_tokens=80_000,
        total_input_tokens=400_000,
        max_calls=12,
    )
    assert len(planned) == 1
    texts = [section["text"] for section in planned[0].payload["sections"]]
    assert texts == ["un", "deux", "trois"]


def test_useful_payload_is_the_ceiling_minus_prompt_and_schema():
    from houston.brain.contract import openai_strict_response_format
    from houston.brain.openai_provider import system_prompt

    assert useful_input_tokens(80_000) == (
        80_000
        - estimate_tokens(system_prompt())
        - estimate_tokens(openai_strict_response_format())
    )


def test_a_section_must_fit_the_useful_payload():
    section = _paragraph("paquet " * 30)
    shaped = dict(section)
    shaped["piece_id"] = "s0000"
    shaped["item_ids"] = []
    cost = estimate_tokens(reading_envelope([shaped], None))
    fixed = provider_fixed_input_tokens()
    planned = plan_lots(
        [section],
        per_call_input_tokens=cost + fixed,
        total_input_tokens=cost * 4,
        max_calls=4,
    )
    assert len(planned) == 1
    with pytest.raises(BrainWorkError) as raised:
        plan_lots(
            [section],
            per_call_input_tokens=cost + fixed - 1,
            total_input_tokens=cost * 4,
            max_calls=4,
        )
    assert raised.value.code == "budget_exceeded"


def test_an_oversized_table_is_split_and_each_lot_repeats_the_header():
    section = _table(["ligne"] * 6)
    one_piece = _table_piece(section, 1)
    one_piece["piece_id"] = "s0000-p0000"
    two_piece = _table_piece(section, 2)
    two_piece["piece_id"] = "s0000"
    one = estimate_tokens(reading_envelope([one_piece], None))
    two = estimate_tokens(reading_envelope([two_piece], None))
    per_call = max(one, two - 1) + provider_fixed_input_tokens()
    planned = plan_lots(
        [section],
        per_call_input_tokens=per_call,
        total_input_tokens=per_call * 20,
        max_calls=12,
    )
    assert len(planned) > 1
    headers = [
        part["header"] == "Mois" and part["unit"] == "EUR"
        for lot in planned
        for part in lot.payload["sections"]
    ]
    parents = [
        part["locator"]["sheet"] == "Suivi"
        for lot in planned
        for part in lot.payload["sections"]
    ]
    assert headers
    assert all(headers)
    assert all(parents)


def test_one_row_above_the_call_ceiling_fails_before_any_lot():
    section = _table(["x" * 400])
    with pytest.raises(BrainWorkError) as raised:
        plan_lots(
            [section],
            per_call_input_tokens=1,
            total_input_tokens=400_000,
            max_calls=12,
        )
    assert raised.value.code == "budget_exceeded"


def test_reserved_lots_above_the_total_fail():
    sections = [_paragraph(f"section {index} " * 30) for index in range(4)]
    shaped = dict(sections[0])
    shaped["piece_id"] = "s0000"
    one = estimate_tokens({"sections": [shaped]})
    with pytest.raises(BrainWorkError) as raised:
        plan_lots(
            sections,
            per_call_input_tokens=one + provider_fixed_input_tokens(),
            total_input_tokens=one,
            max_calls=12,
        )
    assert raised.value.code == "budget_exceeded"


def test_extracted_docx_paragraph_and_xlsx_text_are_in_the_lot_payload():
    from docx import Document
    from openpyxl import Workbook

    document = Document()
    document.add_paragraph("Le service du soir a été lent.")
    docx_buffer = io.BytesIO()
    document.save(docx_buffer)
    docx = extract_document(docx_buffer.getvalue(), "docx")
    book = Workbook()
    sheet = book.active
    sheet.title = "Mars"
    sheet["A1"] = "Commentaire salle"
    xlsx_buffer = io.BytesIO()
    book.save(xlsx_buffer)
    xlsx = extract_document(xlsx_buffer.getvalue(), "xlsx")

    docx_lots = plan_lots(
        sections_from_payload(docx.as_payload()),
        per_call_input_tokens=80_000,
        total_input_tokens=400_000,
        max_calls=12,
    )
    xlsx_lots = plan_lots(
        sections_from_payload(xlsx.as_payload()),
        per_call_input_tokens=80_000,
        total_input_tokens=400_000,
        max_calls=12,
    )
    assert "Le service du soir a été lent." in _texts(docx_lots)
    assert "Commentaire salle" in _texts(xlsx_lots)


def _table_piece(section: dict, count: int) -> dict:
    return {
        "kind": "table",
        "header": section["header"],
        "unit": section["unit"],
        "locator": section["locator"],
        "rows": section["rows"][:count],
    }


def _texts(planned) -> str:
    chunks: list[str] = []
    for lot in planned:
        for section in lot.payload["sections"]:
            if section.get("text"):
                chunks.append(section["text"])
            for row in section.get("rows") or []:
                chunks.append(row.get("text") or "")
    return "\n".join(chunks)


def test_a_measure_table_keeps_one_item_per_row_in_order():
    payload = {
        "qualitative": [],
        "measures": [
            {
                "label": "CA",
                "value": "10",
                "unit": "EUR",
                "raw": "10 EUR",
                "value_kind": "number",
                "definition": "Chiffre encaissé",
                "ambiguities": ["hors taxes"],
                "locator": {"sheet": "Suivi", "header": "Mois", "row": str(index + 2)},
            }
            for index in range(3)
        ],
    }
    section = sections_from_payload(payload)[0]
    assert section["item_ids"] == [
        "document:measure:0000",
        "document:measure:0001",
        "document:measure:0002",
    ]
    assert [row["stable_key"] for row in section["rows"]] == section["item_ids"]
    assert section["rows"][0]["definition"] == "Chiffre encaissé"
    assert section["rows"][0]["ambiguities"] == ["hors taxes"]
    assert "definition" not in section["rows"][0]["text"]
    assert section["rows"][0]["text"] == "CA 10 EUR 10 EUR"


def test_minimal_output_accepts_478_short_ids_and_refuses_the_next():
    from django.conf import settings

    from houston.brain.contract import minimal_call_output

    ceiling = int(settings.HOUSTON_BRAIN_PER_CALL_OUTPUT_TOKENS)
    short = [f"s{index:04d}" for index in range(478)]
    longer = [f"s{index:04d}" for index in range(479)]
    subdivided = [f"s{index:04d}-p0000" for index in range(467)]
    assert estimate_tokens(minimal_call_output(short)) <= ceiling
    assert estimate_tokens(minimal_call_output(longer)) > ceiling
    assert estimate_tokens(minimal_call_output(subdivided)) > ceiling


def test_synthesis_majorant_names_the_last_accepted_lot_count():
    from houston.brain.configuration import empty_knowledge, useful_input_tokens
    from houston.brain.contract import MAX_SUBJECTS, synthesis_input_majorant

    pieces = [f"s{index:04d}" for index in range(467)]
    knowledge = empty_knowledge()
    useful = useful_input_tokens(80_000)
    output_tokens = 8_000

    def bound(lot_count: int) -> int:
        return synthesis_input_majorant(
            pieces,
            key_slots=MAX_SUBJECTS,
            knowledge=knowledge,
            per_call_output_tokens=output_tokens,
            lot_count=lot_count,
        )

    accepted = max(lot_count for lot_count in range(0, 13) if bound(lot_count) <= useful)
    assert bound(accepted) <= useful
    assert bound(accepted + 1) > useful


def test_document_keys_keep_source_indexes_in_capture_and_lots():
    from houston.brain.collection import _document_rows

    payload = {
        "qualitative": [
            {"text": "  "},
            {"text": "Café — déjà"},
            "ignore",
            {"text": "Suite"},
        ],
        "measures": [
            "ignore",
            {
                "label": "CA",
                "value": "1",
                "unit": "EUR",
                "raw": "1",
                "locator": {"sheet": "Suivi"},
            },
        ],
    }
    row_keys = [row.stable_key for row in _document_rows(payload)]
    section_keys = []
    for section in sections_from_payload(payload):
        section_keys.extend(section["item_ids"])
    assert row_keys == [
        "document:qualitative:0001",
        "document:qualitative:0003",
        "document:measure:0000",
    ]
    assert section_keys == row_keys


def test_incremental_envelope_bytes_match_canonical_json():
    knowledge = {
        "formulation": 'déjà "vu" — 日本語 \\ chemin',
        "limits": "ligne\nsuivante\t\u0001",
        "uncertainty": "",
    }
    searches = ["unité", 'code "spécial"', ""]
    sections = [
        {
            "kind": "paragraph",
            "text": "",
            "locator": {},
            "item_ids": [],
            "piece_id": 's"0000"',
        },
        {
            "kind": "table",
            "header": "Mois",
            "unit": None,
            "locator": {"sheet": "Suivi / été"},
            "rows": [
                {
                    "text": "Café ☕",
                    "locator": {"row": 2},
                    "stable_key": "document:measure:0000",
                }
            ],
            "item_ids": ["document:measure:0000"],
            "piece_id": "s0001-p0000",
        },
        {
            "kind": "paragraph",
            "text": "sans piece",
            "locator": {"block": 1},
        },
    ]
    items = [
        {
            "item_id": "signal-1",
            "stable_key": "signal:1",
            "category": "signal",
            "role": "copie",
            "payload": {"text": "preuve — \u2603", "tags": [], "note": None},
        },
        {
            "item_id": "plan-1",
            "stable_key": "plan:1",
            "category": "plan",
            "role": "copie",
            "payload": {"execution_ids": ["ex-1"], "label": "Plan \\ déjà"},
        },
    ]
    meter = _ReadingMeter(knowledge, searches)
    merged = meter.section_group(sections[0])
    for section in sections[1:]:
        merged = _merge(merged, meter.section_group(section))
    merged = _merge(merged, meter.item_group(items[:1]))
    merged = _merge(merged, meter.item_group(items[1:]))
    envelope = reading_envelope(
        merged.sections,
        knowledge,
        searches=searches,
        items=merged.items,
    )
    raw = canonical_json(envelope).encode("utf-8")
    assert meter.payload_bytes(merged.cost) == len(raw)
    assert meter.tokens(merged.cost) == estimate_tokens(envelope)

    bare = _ReadingMeter(None, [])
    one = bare.section_group({"kind": "paragraph", "text": "seul", "piece_id": "s0000"})
    alone = reading_envelope(one.sections, None, searches=[], items=[])
    assert bare.payload_bytes(one.cost) == len(canonical_json(alone).encode("utf-8"))
    assert "items" not in alone
    assert bare.tokens(one.cost) == estimate_tokens(alone)


def test_packed_lot_boundaries_and_hashes_match_the_full_envelope():
    knowledge = {"formulation": "Knowledge été — 中文", "limits": "aucune"}
    searches = ["chevauchement_signals", "unité"]
    sections = [
        _paragraph(f"Café {index} — ligne déjà vue " + ("漢" * (index % 7)) + ' "suite"')
        for index in range(48)
    ]
    longest = max(sections, key=lambda section: len(section["text"]))
    alone = plan_lots(
        [longest],
        per_call_input_tokens=80_000,
        total_input_tokens=400_000,
        max_calls=2,
        knowledge=knowledge,
        searches=searches,
    )
    per_call = alone[0].reserved_input_tokens + provider_fixed_input_tokens()
    useful = useful_input_tokens(per_call)
    planned = plan_lots(
        sections,
        per_call_input_tokens=per_call,
        total_input_tokens=per_call * len(sections),
        max_calls=len(sections),
        knowledge=knowledge,
        searches=searches,
    )
    assert len(planned) > 1
    _assert_lot_boundaries(planned, knowledge, searches, useful)

    row_texts = [f"CA été {index} — {index},00 €" for index in range(120)]
    one_row = plan_lots(
        [_table([max(row_texts, key=len)])],
        per_call_input_tokens=80_000,
        total_input_tokens=400_000,
        max_calls=2,
        knowledge=knowledge,
        searches=searches,
    )
    # A subdivided row carries a longer piece id than the unsplit one-row table.
    table_call = max(
        per_call,
        one_row[0].reserved_input_tokens + provider_fixed_input_tokens() + 16,
    )
    table = _table(row_texts)
    table_plan = plan_lots(
        [table],
        per_call_input_tokens=table_call,
        total_input_tokens=table_call * 120,
        max_calls=120,
        knowledge=knowledge,
        searches=searches,
    )
    assert len(table_plan) > 1
    _assert_lot_boundaries(table_plan, knowledge, searches, useful_input_tokens(table_call))


def _assert_lot_boundaries(planned, knowledge, searches, useful: int) -> None:
    for lot in planned:
        assert lot.input_hash == digest(lot.payload)
        assert lot.reserved_input_tokens == estimate_tokens(lot.payload)
        assert lot.reserved_input_tokens <= useful
    for left, right in zip(planned, planned[1:]):
        right_sections = right.payload["sections"]
        if right_sections:
            merged = reading_envelope(
                [*left.payload["sections"], right_sections[0]],
                knowledge,
                searches=searches,
                items=left.payload.get("items") or None,
            )
        else:
            right_items = right.payload.get("items") or []
            merged = reading_envelope(
                left.payload["sections"],
                knowledge,
                searches=searches,
                items=[*(left.payload.get("items") or []), right_items[0]],
            )
        assert estimate_tokens(merged) > useful
